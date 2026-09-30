"""
drf-excel renders a list as a spreadsheet. Its renderer asks the view for its
serializer while it renders, and its mixin overrides ``finalize_response``.
"""

import io

from django.test import override_settings
from django.urls import path
from drf_excel.mixins import XLSXFileMixin
from drf_excel.renderers import XLSXRenderer
from openpyxl import load_workbook
from rest_framework import viewsets as drf_viewsets
from rest_framework.permissions import AllowAny

from aiodrf import viewsets
from aiodrf.test import count_hops
from tests.base import both_transports
from tests.testapp.models import Author
from tests.testapp.serializers import AuthorSerializer


class Policies(XLSXFileMixin):
    queryset = Author.objects.order_by("pk")
    serializer_class = AuthorSerializer
    authentication_classes = []
    permission_classes = [AllowAny]
    renderer_classes = [XLSXRenderer]
    filename = "authors.xlsx"


class DRFAuthors(Policies, drf_viewsets.ReadOnlyModelViewSet):
    pass


class Authors(Policies, viewsets.ReadOnlyModelViewSet):
    pass


urlpatterns = [
    path("drf/", DRFAuthors.as_view({"get": "list"})),
    path("aiodrf/", Authors.as_view({"get": "list"})),
]


def rows(response):
    sheet = load_workbook(io.BytesIO(response.content)).active
    return [[cell.value for cell in row] for row in sheet.iter_rows()]


@both_transports
class _ExcelTests:
    @classmethod
    def setUpTestData(cls):
        Author.objects.create(name="Ursula")
        Author.objects.create(name="Octavia")

    @override_settings(ROOT_URLCONF=__name__)
    async def test_aiodrf_renders_the_same_workbook(self):
        drf = await self.api("get", "/drf/")
        with count_hops() as hops:
            aiodrf = await self.api("get", "/aiodrf/")
        assert aiodrf.status_code == 200
        assert aiodrf["content-disposition"] == drf["content-disposition"]
        assert "authors.xlsx" in aiodrf["content-disposition"]
        assert rows(aiodrf) == rows(drf)
        assert [row[-1] for row in rows(aiodrf)[1:]] == ["Ursula", "Octavia"]
        if self.transport == "asgi":
            # The list, and the mixin's ``finalize_response``: code written
            # for DRF runs in a thread.
            assert hops.calls == [
                "ListModelMixin._list",
                "XLSXFileMixin.finalize_response",
            ]
